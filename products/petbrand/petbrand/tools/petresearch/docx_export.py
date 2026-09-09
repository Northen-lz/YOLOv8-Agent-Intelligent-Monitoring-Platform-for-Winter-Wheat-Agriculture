# -*- coding: utf-8 -*-
"""
docx_export —— 把「全案汇总」导出为 Word(.docx)（确定性，无 LLM）。

依赖：python-docx（未安装时 build_case_docx 返回 None，由 UI 提示）。
用法：build_case_docx(case_id) -> docx 路径
先确保 outputs/decks/<case_id>/全案汇总.md 存在（可由 DeckBuilderTool 生成）。
"""

import os
import re

from ...core.config import Config

try:
    from docx import Document  # type: ignore
    HAS_DOCX = True
except Exception:  # pragma: no cover
    Document = None
    HAS_DOCX = False


def _plain(s: str) -> str:
    """轻量去 markdown 记号：**加粗**、`代码`、行首 > 引用、[*] 链接标记。"""
    s = s.strip()
    if s.startswith(">"):
        s = s.lstrip("> ").strip()
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    s = s.replace("~~", "")
    return s.strip()


def _fill(doc, md: str):
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(_plain(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_plain(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_plain(line[2:]), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(_plain(re.sub(r"^\s*[-*]\s+", "", line)), style="List Bullet")
        elif re.match(r"^\s*\d+[.、)]\s+", line):
            doc.add_paragraph(_plain(re.sub(r"^\s*\d+[.、)]\s+", "", line)), style="List Number")
        else:
            doc.add_paragraph(_plain(line))


def build_case_docx(case_id: str):
    """把该案的 全案汇总.md 转成 .docx；缺 md 或依赖时返回 None。"""
    if not HAS_DOCX:
        return None
    deck_dir = os.path.join(Config.DECK_DIR, case_id)
    md_path = os.path.join(deck_dir, "全案汇总.md")
    if not os.path.exists(md_path):
        return None
    doc = Document()
    # 基础样式：中文字体
    try:
        from docx.oxml.ns import qn  # type: ignore
        st = doc.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass
    with open(md_path, encoding="utf-8") as f:
        md = f.read()
    _fill(doc, md)
    out = os.path.join(deck_dir, "全案汇总.docx")
    doc.save(out)
    return out
