# -*- coding: utf-8 -*-
"""
ReportGenerationTool —— 农业智能监测报告生成工具

支持三格式：Markdown（零依赖）/ Word（python-docx）/ PDF（reportlab）。
docx/pdf 库缺失时优雅降级，仅生成 md 并在返回中提示。
报告五节：检测区域信息 / 植株检测结果 / 干旱预测结果 / 风险分析 / 农业管理建议
"""

import json
import os
from datetime import datetime
from typing import Any, Dict, List

from ...core.config import Config
from ha_framework.tools.base import BaseTool, ToolParameter


class ReportGenerationTool(BaseTool):
    """监测报告生成工具（md/docx/pdf）"""

    name = "report_generation"
    description = ("生成冬小麦智能监测报告（Markdown/Word/PDF 三格式）。"
                   "用法: report_generation(data={...}, format=\"md|docx|pdf\", filename=\"报告名\")，"
                   "data 含 image/detection/drought/risk_analysis/suggestions 字段。")

    def __init__(self, output_dir: str = None, name: str = None, description: str = None):
        super().__init__(name=name, description=description)
        self.output_dir = output_dir or Config.REPORT_OUTPUT_DIR

    # ---------------- 报告构建 ----------------

    @staticmethod
    def build_markdown(data: Dict[str, Any], title: str = "冬小麦智能监测报告") -> str:
        """从结构化 data 构建 Markdown 报告"""
        md = [f"# {title}", ""]
        md.append(f"> 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  ")
        if data.get("image_path"):
            md.append(f"> 检测图像: {data['image_path']}")
        md.append("")

        # 1. 检测区域信息
        md.append("## 1. 检测区域信息")
        img = data.get("image") or {}
        md.append(f"- 图像尺寸: {img.get('width', '-')} × {img.get('height', '-')} 像素")
        md.append(f"- 检测株数: {data.get('detection', {}).get('count', '-')} 株")
        md.append("")

        # 2. 植株检测结果
        md.append("## 2. 植株检测结果")
        det = data.get("detection") or {}
        md.append(f"- 检测到冬小麦: **{det.get('count', 0)}** 株")
        md.append(f"- 平均置信度: **{det.get('avg_conf', '-')}**")
        md.append(f"- 类别分布: {det.get('class_counts', '-')}")
        for k in ("density_per_10k_px", "uniformity", "coverage"):
            if k in det:
                md.append(f"- {ReportGenerationTool._STAT_LABEL[k]}: {det[k]}")
        md.append("")

        # 3. 干旱预测结果
        md.append("## 3. 干旱预测结果")
        dr = data.get("drought") or {}
        md.append(f"- 已分类植株: {dr.get('total_classified', 0)} 株")
        md.append(f"- 干旱: **{dr.get('drought_count', 0)}** 株 / 正常: **{dr.get('control_count', 0)}** 株")
        md.append(f"- 干旱率: **{dr.get('drought_rate', 0)}**")
        md.append("")

        # 4. 风险分析
        md.append("## 4. 风险分析")
        risk = data.get("risk_analysis") or ""
        md.append(str(risk) if risk else "待分析。")
        md.append("")

        # 5. 农业管理建议
        md.append("## 5. 农业管理建议")
        suggestions = data.get("suggestions") or ""
        md.append(str(suggestions) if suggestions else "待补充。")
        return "\n".join(md)

    _STAT_LABEL = {
        "density_per_10k_px": "植株密度(株/万像素)",
        "uniformity": "空间均匀度",
        "coverage": "覆盖率",
    }

    # ---------------- 格式转换 ----------------

    @staticmethod
    def _md_to_docx(md_text: str, out_path: str) -> bool:
        """Markdown → docx（粗粒度：标题/列表/段落）"""
        try:
            import re
            from docx import Document
        except ImportError:
            return False
        doc = Document()
        doc.add_heading("冬小麦智能监测报告", 0)
        for line in md_text.splitlines():
            line = line.rstrip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=1)
            elif line.startswith("# "):
                continue
            elif line.startswith("- "):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(line.strip())
        doc.save(out_path)
        return True

    @staticmethod
    def _md_to_pdf(md_text: str, out_path: str) -> bool:
        """Markdown → PDF（reportlab platypus 逐行渲染）"""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            return False
        styles = getSampleStyleSheet()
        h1, h2, body = styles["Heading1"], styles["Heading2"], styles["BodyText"]
        body.spaceAfter = 3 * mm
        story = []
        story.append(Paragraph("冬小麦智能监测报告", h1))
        story.append(Spacer(1, 4 * mm))
        for line in md_text.splitlines():
            line = line.rstrip()
            if line.startswith("## "):
                story.append(Paragraph(line[3:].strip(), h2))
            elif line.startswith("- "):
                story.append(Paragraph("• " + line[2:].strip(), body))
            elif line.strip():
                story.append(Paragraph(line.strip(), body))
        doc = SimpleDocTemplate(out_path, pagesize=A4)
        doc.build(story)
        return True

    # ---------------- BaseTool 接口 ----------------

    def get_parameters(self) -> List[ToolParameter]:
        return [
            ToolParameter("data", "string", "报告数据（JSON 字符串或 dict）", required=True),
            ToolParameter("format", "string", "md / docx / pdf", required=False, default="md"),
            ToolParameter("filename", "string", "报告文件名（不含后缀）", required=False),
            ToolParameter("title", "string", "报告标题", required=False),
        ]

    def run(self, *args, **kwargs) -> str:
        if not args and "input" in kwargs:
            raw = kwargs.pop("input")
            if isinstance(raw, dict):
                kwargs.update(raw)
            else:
                kwargs["data"] = raw
        elif len(args) >= 1 and not kwargs.get("data"):
            kwargs["data"] = args[0]

        data = kwargs.get("data")
        if isinstance(data, str):
            try:
                data = json.loads(data)
            except (json.JSONDecodeError, TypeError):
                return "❌ report_generation 失败: data 需为 JSON 字符串或 dict"
        if not isinstance(data, dict):
            return "❌ report_generation 失败: data 需为 dict"
        data.setdefault("image", {})
        data.setdefault("detection", {})
        data.setdefault("drought", {})

        fmt = str(kwargs.get("format") or "md").lower()
        title = str(kwargs.get("title") or "冬小麦智能监测报告")
        filename = str(kwargs.get("filename") or f"wheat_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}")

        os.makedirs(self.output_dir, exist_ok=True)
        md_text = self.build_markdown(data, title)
        md_path = os.path.join(self.output_dir, filename + ".md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_text)

        generated = ["md"]
        notices = []
        if fmt in ("docx", "all"):
            p = os.path.join(self.output_dir, filename + ".docx")
            if self._md_to_docx(md_text, p):
                generated.append("docx")
            else:
                notices.append("python-docx 未安装，跳过 Word 生成")
        if fmt in ("pdf", "all"):
            p = os.path.join(self.output_dir, filename + ".pdf")
            if self._md_to_pdf(md_text, p):
                generated.append("pdf")
            else:
                notices.append("reportlab 未安装，跳过 PDF 生成")

        return json.dumps({
            "title": title,
            "formats_generated": generated,
            "markdown": md_text,
            "paths": {g: os.path.join(self.output_dir, f"{filename}.{g}") for g in generated},
            "notices": notices,
        }, ensure_ascii=False)
